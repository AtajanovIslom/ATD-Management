"""
"Xodim hisobotlari" — xodim faoliyati bo'yicha jamlangan hisobot.

Hisobot qo'lda yozilmaydi: xodim vazifa yoki loyiha bosqichi bo'yicha ishini
topshirgan zahoti u shu jadvalga o'zi tushadi. Xodim xohlasa qo'shimcha
matnli hisobot ham yozadi (work_logs jadvali) — u ham shu yerda chiqadi.

Manbalar (har biri — jadvaldagi bitta qator):
  task     — vazifa bo'yicha topshirilgan hisobot (task_reports)
  project  — loyiha bosqichi bo'yicha topshirilgan hisobot (daily_reports)
  note     — xodimning qo'shimcha hisoboti (work_logs)

Kim kimni ko'radi (utils.scope_user_ids):
  superadmin/direksiya — barcha xodimlar
  boshqarma rahbari    — o'z boshqarmasi xodimlari
  bo'lim rahbari       — o'z bo'limi xodimlari
  oddiy xodim          — faqat o'zi (boshqalarning hisobotini ko'rmaydi)

Endpoint'lar:
  GET /api/employee-reports/employees  — tanlash ro'yxati (bo'ysunuvchilar)
  GET /api/employee-reports            — hisobot (JSON)
  GET /api/employee-reports/export     — PDF yoki Word (?format=pdf|docx)

Umumiy parametrlar:
  period   = today | week | month | custom   (default: today)
  date     = YYYY-MM-DD — davr qaysi kunga nisbatan (default: bugun)
  from/to  = YYYY-MM-DD — period=custom uchun
  user_id  = <id> | all  (default: rahbar uchun all, xodim uchun o'zi)
  division_id / department_id — "Barcha xodimlar" ni toraytirish
  lang     = uz | ru  — eksport tili
"""
import io
import os
import re
from datetime import date as _date, datetime, time, timedelta, timezone
from xml.sax.saxutils import escape

from flask import Blueprint, jsonify, request, send_file
from flask_jwt_extended import get_jwt, get_jwt_identity, jwt_required

from app import db
from app.models import (
    DailyReport, Department, Division, Project, ProjectStage, Task, TaskReport,
    User, WorkLog,
)
from app.utils import get_scope, is_any_admin, scope_user_ids, FULL_ACCESS_ROLES

employee_reports_bp = Blueprint('employee_reports', __name__)

PERIODS = ('today', 'week', 'month', 'custom')
MAX_CUSTOM_DAYS = 366
# Bitta katakdagi matn chegarasi — PDF jadval qatori sahifadan uzun
# bo'lib qolsa reportlab uni joylashtira olmaydi
MAX_CELL_CHARS = 1500

# Eksport matnlari (sahifa tili bo'yicha)
TEXTS = {
    'uz': {
        'title': 'Xodimlar faoliyati hisoboti',
        'title_one': 'Xodim faoliyati hisoboti',
        'period': 'Davr',
        'generated': 'Shakllantirildi',
        'by': 'Shakllantirdi',
        'employees': 'Xodimlar soni',
        'total': 'Jami hisobotlar',
        'col_no': '№',
        'col_kind': 'Turi',
        'col_title': 'Vazifa / ish nomi',
        'col_time': 'Topshirilgan vaqti',
        'col_status': 'Holati',
        'col_content': 'Hisobot matni',
        'kind_task': 'Vazifa',
        'kind_project': 'Loyiha',
        'kind_note': "Qo'shimcha",
        'empty': "Ushbu davrda hisobot topshirilmagan.",
        'summary': 'Jami: {total} ta (vazifa: {task}, loyiha: {project}, qo\'shimcha: {note})',
        'signature': 'Imzo',
        'page': 'Sahifa',
        'no_division': "Bo'lim ko'rsatilmagan",
        'status': {
            'active': 'Faol', 'in_progress': 'Jarayonda', 'review': 'Tekshiruvda',
            'returned': 'Qaytarilgan', 'completed': 'Tasdiqlangan', 'cancelled': 'Bekor qilingan',
            'pending': 'Kutilmoqda', 'approved': 'Tasdiqlangan', 'note': '—',
        },
    },
    'ru': {
        'title': 'Отчёт о деятельности сотрудников',
        'title_one': 'Отчёт о деятельности сотрудника',
        'period': 'Период',
        'generated': 'Сформирован',
        'by': 'Сформировал',
        'employees': 'Количество сотрудников',
        'total': 'Всего отчётов',
        'col_no': '№',
        'col_kind': 'Тип',
        'col_title': 'Задача / работа',
        'col_time': 'Время сдачи',
        'col_status': 'Статус',
        'col_content': 'Текст отчёта',
        'kind_task': 'Задача',
        'kind_project': 'Проект',
        'kind_note': 'Дополнительный',
        'empty': 'За этот период отчёты не сдавались.',
        'summary': 'Всего: {total} (задачи: {task}, проекты: {project}, доп.: {note})',
        'signature': 'Подпись',
        'page': 'Стр.',
        'no_division': 'Отдел не указан',
        'status': {
            'active': 'Активна', 'in_progress': 'В работе', 'review': 'На проверке',
            'returned': 'Возвращена', 'completed': 'Принята', 'cancelled': 'Отменена',
            'pending': 'Ожидает', 'approved': 'Принята', 'note': '—',
        },
    },
}


# =========================================================================
# DAVR VA VAQT
# =========================================================================

def _parse_date(s):
    if not s:
        return None
    try:
        return _date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return None


def _period_range(args):
    """(key, d_from, d_to) yoki xato matni. Sanalar — mahalliy (Toshkent)
    kalendar kunlari; konteynerda TZ=Asia/Tashkent o'rnatilgan."""
    key = (args.get('period') or 'today').strip()
    if key not in PERIODS:
        return None, "Noto'g'ri davr"
    anchor = _parse_date(args.get('date')) or _date.today()
    if key == 'today':
        return (key, anchor, anchor), None
    if key == 'week':
        start = anchor - timedelta(days=anchor.weekday())  # dushanba
        return (key, start, start + timedelta(days=6)), None
    if key == 'month':
        start = anchor.replace(day=1)
        nxt = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
        return (key, start, nxt - timedelta(days=1)), None
    d_from = _parse_date(args.get('from'))
    d_to = _parse_date(args.get('to'))
    if not d_from or not d_to:
        return None, "Davr boshi va oxiri kiritilishi shart"
    if d_from > d_to:
        d_from, d_to = d_to, d_from
    if (d_to - d_from).days > MAX_CUSTOM_DAYS:
        return None, "Davr bir yildan oshmasligi kerak"
    return (key, d_from, d_to), None


def _utc_bounds(d_from, d_to):
    """Mahalliy kunlar oralig'ini bazadagi naive UTC vaqt oralig'iga o'giradi.
    Ilgari mahalliy sana UTC vaqt bilan to'g'ridan-to'g'ri solishtirilardi —
    tungi 00:00–05:00 orasidagi ishlar kechagi kunga tushib qolardi."""
    start = datetime.combine(d_from, time.min).astimezone(timezone.utc).replace(tzinfo=None)
    end = datetime.combine(d_to + timedelta(days=1), time.min).astimezone(timezone.utc).replace(tzinfo=None)
    return start, end


def _as_utc(dt):
    if dt is None:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def _local(dt):
    dt = _as_utc(dt)
    return dt.astimezone() if dt else None


def _fmt_dt(dt):
    loc = _local(dt)
    return loc.strftime('%d.%m.%Y %H:%M') if loc else ''


# =========================================================================
# KIMNING HISOBOTI
# =========================================================================

def _viewer():
    role, dept_id, div_id = get_scope(get_jwt())
    me = int(get_jwt_identity())
    leader = is_any_admin(role)
    scope = scope_user_ids(role, dept_id, div_id, me) if leader else {me}
    return role, me, leader, scope


def _employee_query(scope):
    q = User.query.filter(User.is_active.is_(True), User.role != 'agent')
    if scope is not None:
        q = q.filter(User.id.in_(scope)) if scope else q.filter(db.false())
    return q


def _target_users(args):
    """Hisobotga kiradigan xodimlar yoki (None, xato, status)."""
    role, me, leader, scope = _viewer()
    raw = (args.get('user_id') or '').strip()

    if not leader:
        # Oddiy xodim faqat o'z hisobotini ko'radi va yuklaydi
        if raw and raw not in ('all', str(me)):
            return None, "Boshqa xodimning hisobotini ko'rish huquqingiz yo'q", 403
        return [User.query.get(me)], None, 200

    q = _employee_query(scope)
    if raw and raw != 'all':
        try:
            uid = int(raw)
        except ValueError:
            return None, "user_id noto'g'ri", 400
        user = q.filter(User.id == uid).first()
        if user is None:
            return None, "Bu xodim sizning ko'lamingizda emas", 403
        return [user], None, 200

    div_id = args.get('division_id', type=int)
    dept_id = args.get('department_id', type=int)
    if div_id:
        q = q.filter(User.division_id == div_id)
    if dept_id and role in FULL_ACCESS_ROLES:
        div_ids = db.session.query(Division.id).filter(Division.department_id == dept_id)
        q = q.filter(db.or_(User.department_id == dept_id, User.division_id.in_(div_ids)))
    return q.order_by(User.full_name).all(), None, 200


# =========================================================================
# MA'LUMOT YIG'ISH
# =========================================================================

def _collect(users, d_from, d_to):
    """{user_id: [item, ...]} — vaqt bo'yicha tartiblangan."""
    ids = [u.id for u in users]
    out = {uid: [] for uid in ids}
    if not ids:
        return out
    start, end = _utc_bounds(d_from, d_to)

    task_reports = TaskReport.query.filter(
        TaskReport.user_id.in_(ids),
        TaskReport.created_at >= start,
        TaskReport.created_at < end,
    ).all()
    task_ids = {r.task_id for r in task_reports}
    tasks = {t.id: t for t in Task.query.filter(Task.id.in_(task_ids)).all()} if task_ids else {}
    for r in task_reports:
        t = tasks.get(r.task_id)
        if t is None:
            continue
        out[r.user_id].append({
            'kind': 'task',
            'id': f'task-{r.id}',
            'ref_id': t.id,
            'title': t.name,
            'at': _as_utc(r.created_at),
            'status': t.status or 'active',
            'approved_at': _as_utc(t.completed_at) if t.status == 'completed' else None,
            'content': r.content or '',
            'files': len(r.files),
        })

    project_reports = DailyReport.query.filter(
        DailyReport.user_id.in_(ids),
        DailyReport.created_at >= start,
        DailyReport.created_at < end,
    ).all()
    project_ids = {r.project_id for r in project_reports}
    projects = {p.id: p for p in Project.query.filter(Project.id.in_(project_ids)).all()} if project_ids else {}
    stage_ids = {r.stage_id for r in project_reports if r.stage_id}
    stages = {s.id: s for s in ProjectStage.query.filter(ProjectStage.id.in_(stage_ids)).all()} if stage_ids else {}
    for r in project_reports:
        p = projects.get(r.project_id)
        if p is None:
            continue
        st = stages.get(r.stage_id)
        out[r.user_id].append({
            'kind': 'project',
            'id': f'project-{r.id}',
            'ref_id': p.id,
            'title': f"{p.name} — {st.name}" if st else p.name,
            'at': _as_utc(r.created_at),
            'status': (st.status if st else p.status) or 'in_progress',
            'approved_at': _as_utc(st.completed_at) if st and st.status == 'completed' else None,
            'content': r.content or '',
            'files': len(r.files),
        })

    notes = WorkLog.query.filter(
        WorkLog.user_id.in_(ids),
        WorkLog.work_date >= d_from,
        WorkLog.work_date <= d_to,
    ).all()
    for w in notes:
        created = _as_utc(w.created_at)
        out[w.user_id].append({
            'kind': 'note',
            'id': f'note-{w.id}',
            'ref_id': w.id,
            'title': None,
            'at': created,
            'work_date': w.work_date,
            'status': 'note',
            'approved_at': None,
            'content': w.content or '',
            'files': 0,
        })

    def sort_key(it):
        if it['kind'] == 'note':
            loc = _local(it['at'])
            return (it['work_date'], loc.time() if loc else time.min)
        loc = _local(it['at'])
        return (loc.date(), loc.time())

    for uid in out:
        out[uid].sort(key=sort_key)
    return out


def _totals(items):
    t = {'task': 0, 'project': 0, 'note': 0}
    for it in items:
        t[it['kind']] += 1
    t['total'] = len(items)
    return t


def _employee_info(u):
    div = u.division
    dept = u.managed_department or (div.department if div else None)
    return {
        'id': u.id,
        'full_name': u.full_name,
        'position': u.position or '',
        'role': u.role,
        'division_name': div.name if div else None,
        'department_name': dept.name if dept else None,
    }


def _build(args):
    period, err = _period_range(args)
    if err:
        return None, err, 400
    users, err, status = _target_users(args)
    if err:
        return None, err, status
    key, d_from, d_to = period
    data = _collect(users, d_from, d_to)
    employees = []
    for u in users:
        items = data.get(u.id, [])
        employees.append({**_employee_info(u), 'items': items, 'totals': _totals(items)})
    all_items = [it for e in employees for it in e['items']]
    return {
        'period': {'key': key, 'from': d_from, 'to': d_to},
        'employees': employees,
        'totals': {**_totals(all_items), 'employees': len(employees)},
    }, None, 200


def _item_json(it):
    return {
        'kind': it['kind'],
        'id': it['id'],
        'ref_id': it['ref_id'],
        'title': it['title'],
        'at': it['at'].isoformat() if it['at'] else None,
        'work_date': it['work_date'].isoformat() if it.get('work_date') else None,
        'status': it['status'],
        'approved_at': it['approved_at'].isoformat() if it['approved_at'] else None,
        'content': it['content'],
        'files': it['files'],
    }


# =========================================================================
# ENDPOINT'LAR
# =========================================================================

@employee_reports_bp.route('/employees', methods=['GET'])
@jwt_required()
def list_employees():
    """Hisobot uchun tanlanadigan xodimlar. Oddiy xodimga — faqat o'zi."""
    role, me, leader, scope = _viewer()
    users = _employee_query(scope).order_by(User.full_name).all()
    divisions = {}
    for u in users:
        if u.division:
            divisions[u.division.id] = u.division.name
    departments = []
    if role in FULL_ACCESS_ROLES:
        departments = [{'id': d.id, 'name': d.name}
                       for d in Department.query.order_by(Department.name).all()]
    return jsonify({
        'is_leader': leader,
        'can_filter_department': role in FULL_ACCESS_ROLES,
        'employees': [_employee_info(u) for u in users],
        'divisions': [{'id': k, 'name': v} for k, v in sorted(divisions.items(), key=lambda x: x[1])],
        'departments': departments,
    })


@employee_reports_bp.route('', methods=['GET'])
@jwt_required()
def get_report():
    report, err, status = _build(request.args)
    if err:
        return jsonify({'error': err}), status
    return jsonify({
        'period': {
            'key': report['period']['key'],
            'from': report['period']['from'].isoformat(),
            'to': report['period']['to'].isoformat(),
        },
        'totals': report['totals'],
        'employees': [
            {**{k: v for k, v in e.items() if k != 'items'},
             'items': [_item_json(it) for it in e['items']]}
            for e in report['employees']
        ],
    })


@employee_reports_bp.route('/export', methods=['GET'])
@jwt_required()
def export_report():
    fmt = (request.args.get('format') or 'pdf').lower()
    if fmt not in ('pdf', 'docx'):
        return jsonify({'error': "Format pdf yoki docx bo'lishi kerak"}), 400
    lang = request.args.get('lang') if request.args.get('lang') in TEXTS else 'uz'

    report, err, status = _build(request.args)
    if err:
        return jsonify({'error': err}), status

    viewer = User.query.get(int(get_jwt_identity()))
    tx = TEXTS[lang]
    d_from, d_to = report['period']['from'], report['period']['to']
    meta = {
        'title': tx['title_one'] if len(report['employees']) == 1 else tx['title'],
        'period': _period_text(d_from, d_to),
        'generated': datetime.now().strftime('%d.%m.%Y %H:%M'),
        'by': viewer.full_name if viewer else '',
    }

    if fmt == 'pdf':
        buf = _render_pdf(report, meta, tx)
        mimetype = 'application/pdf'
    else:
        buf = _render_docx(report, meta, tx)
        mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

    suffix = d_from.isoformat() if d_from == d_to else f"{d_from.isoformat()}_{d_to.isoformat()}"
    fname = f"xodim_hisoboti_{suffix}.{fmt}"
    return send_file(buf, as_attachment=True, download_name=fname, mimetype=mimetype)


def _period_text(d_from, d_to):
    if d_from == d_to:
        return d_from.strftime('%d.%m.%Y')
    return f"{d_from.strftime('%d.%m.%Y')} — {d_to.strftime('%d.%m.%Y')}"


# =========================================================================
# EKSPORT: UMUMIY
# =========================================================================

def _row_values(it, tx):
    """Jadval qatori: (turi, nomi, vaqti, holati, matni)."""
    kind = tx['kind_' + it['kind']]
    if it['kind'] == 'note':
        title = '—'
        when = it['work_date'].strftime('%d.%m.%Y') if it.get('work_date') else ''
        status = tx['status']['note']
    else:
        title = it['title'] or ''
        when = _fmt_dt(it['at'])
        status = tx['status'].get(it['status'], it['status'])
        if it['approved_at']:
            status = f"{status}\n{_fmt_dt(it['approved_at'])}"
    content = (it['content'] or '').strip()
    if len(content) > MAX_CELL_CHARS:
        content = content[:MAX_CELL_CHARS].rstrip() + '…'
    return kind, title, when, status, content


def _employee_heading(e, tx):
    parts = [e['full_name']]
    if e['position']:
        parts.append(e['position'])
    unit = e['division_name'] or e['department_name']
    if unit:
        parts.append(unit)
    return ' — '.join(parts)


def _summary_line(e, tx):
    t = e['totals']
    return tx['summary'].format(total=t['total'], task=t['task'], project=t['project'], note=t['note'])


# =========================================================================
# EKSPORT: PDF (reportlab)
# =========================================================================

_FONT_CANDIDATES = [
    # Docker image'da fonts-dejavu-core paketi o'rnatiladi (kirill harflari uchun)
    ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
     '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
    # Lokal ishlab chiqish (Windows)
    ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf'),
]
_fonts = None


def _pdf_fonts():
    """(oddiy, qalin) shrift nomlari. TTF topilmasa Helvetica — lekin unda
    kirill harflari chiqmaydi, shuning uchun serverda DejaVu shart."""
    global _fonts
    if _fonts:
        return _fonts
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont('ReportSans', regular))
            pdfmetrics.registerFont(TTFont('ReportSans-Bold', bold))
            _fonts = ('ReportSans', 'ReportSans-Bold')
            return _fonts
    _fonts = ('Helvetica', 'Helvetica-Bold')
    return _fonts


# Emoji va boshqa "astral" belgilar PDF shriftida yo'q — kvadrat bo'lib
# chiqmasligi uchun olib tashlanadi (Word'da esa qoladi)
_NON_BMP = re.compile('[\U00010000-\U0010FFFF\uFE0F\u200D\u2600-\u27BF]')


def _pdf_text(s):
    s = _NON_BMP.sub('', s or '')
    return escape(s).replace('\n', '<br/>')


def _render_pdf(report, meta, tx):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    regular, bold = _pdf_fonts()
    page = landscape(A4)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=page,
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
        title=meta['title'], author=meta['by'],
    )

    st_title = ParagraphStyle('t', fontName=bold, fontSize=15, leading=19, alignment=TA_CENTER, spaceAfter=2)
    st_meta = ParagraphStyle('m', fontName=regular, fontSize=9, leading=12, alignment=TA_CENTER,
                             textColor=colors.HexColor('#555555'))
    st_emp = ParagraphStyle('e', fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=2,
                            keepWithNext=1)
    st_sum = ParagraphStyle('s', fontName=regular, fontSize=8.5, leading=11,
                            textColor=colors.HexColor('#555555'), spaceAfter=4, keepWithNext=1)
    st_cell = ParagraphStyle('c', fontName=regular, fontSize=8.5, leading=11)
    st_head = ParagraphStyle('h', fontName=bold, fontSize=8.5, leading=11, textColor=colors.white)
    st_empty = ParagraphStyle('x', fontName=regular, fontSize=9, leading=12,
                              textColor=colors.HexColor('#777777'))

    story = [
        Paragraph(_pdf_text(meta['title']), st_title),
        Paragraph(_pdf_text(f"{tx['period']}: {meta['period']}"), st_meta),
        Paragraph(_pdf_text(
            f"{tx['employees']}: {report['totals']['employees']} · {tx['total']}: {report['totals']['total']}"
            f" · {tx['by']}: {meta['by']} · {tx['generated']}: {meta['generated']}"), st_meta),
        Spacer(1, 6),
    ]

    usable = page[0] - doc.leftMargin - doc.rightMargin
    widths = [9 * mm, 24 * mm, 62 * mm, 30 * mm, 30 * mm]
    widths.append(usable - sum(widths))
    header = [Paragraph(_pdf_text(tx[k]), st_head) for k in
              ('col_no', 'col_kind', 'col_title', 'col_time', 'col_status', 'col_content')]
    tstyle = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4e79')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#b7c4d1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f6f9')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ])

    for e in report['employees']:
        head = [Paragraph(_pdf_text(_employee_heading(e, tx)), st_emp),
                Paragraph(_pdf_text(_summary_line(e, tx)), st_sum)]
        if not e['items']:
            story.append(KeepTogether(head + [Paragraph(_pdf_text(tx['empty']), st_empty)]))
            continue
        rows = [header]
        for i, it in enumerate(e['items'], 1):
            kind, title, when, status, content = _row_values(it, tx)
            rows.append([Paragraph(str(i), st_cell)] + [
                Paragraph(_pdf_text(v), st_cell) for v in (kind, title, when, status, content)])
        table = Table(rows, colWidths=widths, repeatRows=1)
        table.setStyle(tstyle)
        # Qisqa jadval sarlavhasi bilan bir sahifada qolsin; uzun jadval
        # sahifalarga bo'linadi (sarlavha keyingi qator bilan birga ko'chadi)
        if len(rows) <= 8:
            story.append(KeepTogether(head + [table]))
        else:
            story.extend(head + [table])

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor('#777777'))
        canvas.drawRightString(page[0] - doc.rightMargin, 7 * mm,
                               f"{tx['page']} {canvas.getPageNumber()}")
        canvas.drawString(doc.leftMargin, 7 * mm, _NON_BMP.sub('', meta['title']))
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    return buf


# =========================================================================
# EKSPORT: WORD (python-docx)
# =========================================================================

def _render_docx(report, meta, tx):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(section, side, Cm(1.5))

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)

    h = doc.add_heading(meta['title'], level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        f"{tx['period']}: {meta['period']}",
        f"{tx['employees']}: {report['totals']['employees']} · {tx['total']}: {report['totals']['total']}",
        f"{tx['by']}: {meta['by']} · {tx['generated']}: {meta['generated']}",
    ):
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    widths = [Cm(1.0), Cm(2.6), Cm(6.5), Cm(3.0), Cm(3.0), Cm(10.6)]
    headers = [tx[k] for k in ('col_no', 'col_kind', 'col_title', 'col_time', 'col_status', 'col_content')]

    def shade(cell, hex_color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tc_pr.append(shd)

    def repeat_header(row):
        tr_pr = row._tr.get_or_add_trPr()
        el = OxmlElement('w:tblHeader')
        el.set(qn('w:val'), 'true')
        tr_pr.append(el)

    for e in report['employees']:
        doc.add_heading(_employee_heading(e, tx), level=2)
        p = doc.add_paragraph(_summary_line(e, tx))
        p.runs[0].font.size = Pt(9)
        p.runs[0].italic = True
        if not e['items']:
            doc.add_paragraph(tx['empty'])
            continue

        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        hdr = table.rows[0]
        repeat_header(hdr)
        for i, name in enumerate(headers):
            cell = hdr.cells[i]
            cell.width = widths[i]
            cell.text = name
            shade(cell, '1F4E79')
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for n, it in enumerate(e['items'], 1):
            values = (str(n),) + _row_values(it, tx)
            cells = table.add_row().cells
            for i, v in enumerate(values):
                cells[i].width = widths[i]
                cells[i].text = v
                for para in cells[i].paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
        doc.add_paragraph()

    p = doc.add_paragraph(f"{tx['signature']}: ____________________")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
